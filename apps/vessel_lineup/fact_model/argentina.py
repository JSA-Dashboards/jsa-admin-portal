"""Argentina (and Uruguay) from the Nabsa agency e-mail: DAILY VSL LINE UP / SAILED VSLS.

Each daily e-mail carries two workbooks:
  lineupYYYYMMDD.xlsx  sheet 'DailyVesselLineUp' (Port, Terminal, Vessel, ETA, ETB, ETF, Ops, Tons, Commodity,
                       Destination, Origin, Charterer) - a full snapshot of vessels queued, loading AND discharging.
  sailedYYYYMMDD.xlsx  month-to-date list of sailed vessels (Port, Terminal, Vessel, Status, Date, Tons, Cargo,
                       Origin, Destination, Coordinator, Charterer); it restarts each month.

Only EXPORTS count: line-up rows with Ops == 'load', and sailed rows whose Origin is a Plata-basin country
(imports - fertiliser, steel, ... - arrive with a foreign origin). A vessel that loads several cargoes
(corn + meal) appears as one row per cargo; its vessel count is split across those cargoes in proportion to
tonnage so that region totals still equal the number of vessels (same idea as the US combo split).
"""

import re
import warnings

import pandas as pd

from . import events as ev
from . import lineup as lu
from . import store
from .dimensions import add_geography, commodity_group, default_config

LOAD_ORIGINS = {"ARGENTINA", "URUGUAY", "PARAGUAY", "BOLIVIA", "NOT AVAILABLE", ""}


def _clean(v):
    return None if v is None else (str(v).replace("\x00", "").strip() or None)


def _table(path, sheet=None):
    from openpyxl import load_workbook
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        wb = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb[sheet] if sheet else wb.worksheets[0]
        rows = list(ws.iter_rows(values_only=True))
    finally:
        wb.close()
    h = next((i for i, r in enumerate(rows) if r and str(r[0]).strip().lower() == "port"), None)
    if h is None:
        raise ValueError(f"{path}: no 'Port' header row found")
    cols = [str(c).strip() if c is not None else f"c{i}" for i, c in enumerate(rows[h])]
    head = [r for r in rows[:h] if r]
    return head, pd.DataFrame(rows[h + 1:], columns=cols)


def report_date_from_name(name):
    m = re.search(r"(\d{4})(\d{2})(\d{2})", name or "")
    return pd.Timestamp(int(m.group(1)), int(m.group(2)), int(m.group(3))) if m else None


def read_lineup(path):
    """Vessel-level exports in the queue: port_raw, berth, vessel, kmt, commodity_raw, destination, charterer, eta_text."""
    head, d = _table(path, "DailyVesselLineUp")
    rep = next((pd.Timestamp(c) for r in head for c in r if hasattr(c, "year")), None)
    d = d[d["Vessel"].notna() & (d["Vessel"].astype(str).str.strip().str.upper() != "NIL")]
    d = d[d["Ops"].astype(str).str.strip().str.lower() == "load"]
    out = pd.DataFrame({
        "port_raw": d["Port"].map(_clean), "berth": d["Terminal"].map(_clean), "vessel": d["Vessel"].map(_clean),
        "kmt": pd.to_numeric(d["Tons"], errors="coerce") / 1000.0,
        "commodity_raw": d["Commodity"].map(_clean), "destination": d["Destination"].map(_clean),
        "charterer": d["Charterer"].map(_clean),
        "eta_text": (d["ETA"].fillna("").astype(str) + " " + d["ETB"].fillna("").astype(str) + " "
                     + d["ETF"].fillna("").astype(str)).str.strip(),
    }).reset_index(drop=True)
    out.attrs["report_date"] = rep
    return out


def read_sailed(path):
    """Vessel/cargo-level export sailings: port_raw, berth, vessel, sail_date, kmt, commodity_raw, destination, charterer."""
    _, d = _table(path)
    d = d[d["Vessel"].notna()]
    d = d[d["Origin"].map(lambda x: (_clean(x) or "").upper() in LOAD_ORIGINS)]
    return pd.DataFrame({
        "port_raw": d["Port"].map(_clean), "berth": d["Terminal"].map(_clean), "vessel": d["Vessel"].map(_clean),
        "sail_date": pd.to_datetime(d["Date"], errors="coerce"),
        "kmt": pd.to_numeric(d["Tons"], errors="coerce") / 1000.0,
        "commodity_raw": d["Cargo"].map(_clean), "destination": d["Destination"].map(_clean),
        "charterer": d["Charterer"].map(_clean),
    }).dropna(subset=["sail_date"]).reset_index(drop=True)


def _vessel_weights(df, keys):
    """1 per vessel, split across its cargo lines by tonnage (equal split when no tonnage)."""
    k = df["kmt"].fillna(0.0)
    grp = [df[c].fillna("") for c in keys]
    tot = k.groupby(grp).transform("sum")
    n = k.groupby(grp).transform("size")
    return (k / tot).where(tot > 0, 1.0 / n)


def argentina_lineup_snapshot(vessels, report_date, cfg=None):
    cfg = cfg or default_config()
    q = add_geography(vessels.copy(), "Argentina", cfg)
    q["source"] = "Argentina"
    q["report_date"] = pd.Timestamp(report_date)
    q["commodity"] = [commodity_group("Argentina", c) for c in q["commodity_raw"]]
    q["combo"] = ""
    q["vessel"] = q["vessel"].fillna("?")
    q["_w"] = _vessel_weights(q, ["port", "vessel"])
    snap = lu._aggregate(q)
    # _aggregate counts distinct vessels; Argentina needs the pro-rata count so totals equal vessels
    qq = q.assign(destination=q["destination"].astype("string").str.strip().str.upper()
                  .replace({"": "UNKNOWN"}).fillna("UNKNOWN"))
    w = qq.groupby(lu.DIMS, dropna=False)["_w"].sum().reset_index().rename(columns={"_w": "_vw"})
    w["report_date"] = pd.to_datetime(w["report_date"])
    snap = snap.merge(w, on=lu.DIMS, how="left")
    snap["vessels"] = snap["_vw"]
    return snap[lu.FACT_COLS]


def argentina_sail_events(sailed, cfg=None):
    cfg = cfg or default_config()
    if sailed.empty:
        return pd.DataFrame(columns=ev.EVENT_COLS + ["vessel_w"])
    s = add_geography(sailed.copy(), "Argentina", cfg)
    s["commodity"] = [commodity_group("Argentina", c) for c in s["commodity_raw"]]
    s["combo"] = ""
    s["_d"] = s["sail_date"].dt.strftime("%Y-%m-%d")
    # identity = vessel + sail date + cargo + destination: the terminal/port wording is revised between files
    s["event_id"] = ev._event_ids(s, ["vessel", "_d", "commodity_raw", "destination"])
    s["event_id"] = "AR" + s["event_id"].str[:18]
    s["kmt"] = s["kmt"].where(s["kmt"] > 0)
    s["vessel_w"] = _vessel_weights(s, ["vessel", "_d"])
    out = ev._finish(s, "Argentina", cfg)
    out["vessel_w"] = s["vessel_w"].values
    return out


def ingest_argentina_files(lineup_path=None, sailed_path=None, report_date=None, facts_db=None, cfg=None):
    cfg = cfg or default_config()
    res = {}
    if lineup_path:
        v = read_lineup(lineup_path)
        rd = report_date or v.attrs.get("report_date") or report_date_from_name(str(lineup_path))
        snap = argentina_lineup_snapshot(v, rd, cfg)
        store.save_lineup_snapshot(snap, facts_db)
        res.update(report_date=str(pd.Timestamp(rd).date()), queued_kmt=round(float(snap["kmt"].sum()), 1),
                   queued_vessels=round(float(snap["vessels"].sum()), 1))
    if sailed_path:
        e = argentina_sail_events(read_sailed(sailed_path), cfg)
        res["sail_events"] = store.save_events(e, facts_db)
    return res
