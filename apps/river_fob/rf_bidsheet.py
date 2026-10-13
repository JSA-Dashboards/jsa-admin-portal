"""Parse the daily **Bid Sheet** Excel (Doug Schultz's forwarded email) into the
River FOB archive.

The workbook's **'Bid Sheet'** tab, cells A1:T16, carries the barge-freight and
CIF NOLA tables; **A1 is the as-of date**. This module is email-source-agnostic:
hand `save_bidsheet()` a path to the .xlsx and it parses + writes to Snowflake.
The email fetch (Outlook COM on the desktop, or Graph on the droplet) is a thin
front-end that just produces the file, then calls this.

Layout (1-indexed columns on the 'Bid Sheet' tab):
  A        month label (TW/NW are spot rows, skipped; then Sept..Apr)
  B..I     barge freight, already the tariff multiplier (7.75 == 775%):
           B ILL->IL, C OHIO->Ohio, D L OHIO (ignored, dup of OHIO),
           E MM->Davenport South+McGregor South, F CITIES->Upper Miss,
           G STL->STL, H MTCT->Lower Miss, I ARK (ignored, not a FOB region)
  M/N O/P Q/R S/T   CIF value + contract letter for CORN / MILO(ignored) /
           BEANS / WHEAT; CIF is cents -> $/bu (/100).
Blank cells are closed reaches / absent months and are simply skipped.
"""
import datetime as dt
import os
import re

import openpyxl

import rf_db as db
import rf_fob_model as M
from rf_paste_parse import _MONTHS


def date_from_filename(name):
    """MMDDYY in the file name -> date (e.g. '100226.xlsx' -> 2026-10-02), else
    None. Doug names the file by the sheet's REAL date; the A1 cell is NOT
    reliable (it has been seen a day behind the actual sheet)."""
    m = re.search(r"(\d{2})(\d{2})(\d{2})", os.path.basename(str(name or "")))
    if not m:
        return None
    mm, dd, yy = (int(x) for x in m.groups())
    try:
        return dt.date(2000 + yy, mm, dd)
    except ValueError:
        return None

# Freight source column -> FOB freight region(s). D (L OHIO) and I (ARK) omitted.
FREIGHT_COLS = {2: ["IL"], 3: ["Ohio"], 5: ["Davenport South", "McGregor South"],
                6: ["Upper Miss"], 7: ["STL"], 8: ["Lower Miss"]}
# commodity -> (value column, contract-letter column). MILO (O/P) omitted.
CIF_COLS = {"Corn": (13, 14), "Soybeans": (17, 18), "Wheat": (19, 20)}
_PRE = {"Corn": "C", "Soybeans": "S", "Wheat": "W"}


def parse_bidsheet(path, as_of=None):
    """-> (as_of: date, cif, freight, contracts). Raises if there's no Bid Sheet tab.

    The as-of date comes from (in priority): an explicit `as_of`, then the file
    NAME (MMDDYY), then the A1 cell as a last resort. A1 is unreliable — it has
    been a day behind the real sheet — so callers should pass the date derived
    from the attachment filename / email."""
    wb = openpyxl.load_workbook(path, data_only=True)
    if "Bid Sheet" not in wb.sheetnames:
        raise ValueError(f"No 'Bid Sheet' tab in {path} (tabs: {wb.sheetnames})")
    ws = wb["Bid Sheet"]
    if as_of is None:
        a1 = ws["A1"].value
        a1_date = (a1.date() if isinstance(a1, dt.datetime)
                   else (dt.date.fromisoformat(str(a1)[:10]) if a1 else None))
        as_of = date_from_filename(path) or a1_date
    if as_of is None:
        raise ValueError(f"Could not determine as-of date for {path} "
                         "(no override, no date in filename, empty A1).")

    cif, freight, contracts = {}, {}, {}
    for r in range(3, 13):                       # TW, NW, Sept..Apr
        lbl = ws.cell(r, 1).value
        mon = _MONTHS.get(str(lbl).strip().upper()) if lbl else None
        if not mon:                              # TW / NW spot rows
            continue
        for com, (vc, lc) in CIF_COLS.items():
            v = ws.cell(r, vc).value
            if isinstance(v, (int, float)):
                cif.setdefault(com, {})[mon] = round(v / 100.0, 4)
                lt = ws.cell(r, lc).value
                if lt:
                    contracts.setdefault(com, {})[mon] = str(lt).strip().upper()
        for col, regs in FREIGHT_COLS.items():
            v = ws.cell(r, col).value
            if isinstance(v, (int, float)):      # already the multiplier (no /100)
                for reg in regs:
                    freight.setdefault(reg, {})[mon] = round(float(v), 4)
    return as_of, cif, freight, contracts


def build_payload(as_of, cif, freight, contracts):
    """Roll the delivery window/contract chain to the as-of date (like the app),
    override the computed contract with the sheet's letters, and clip CIF/freight
    to the window. -> (cif, freight, calendar) ready for db.save_snapshot."""
    months = M.months_for(as_of)
    cal = {}
    for com in M.COMMODITIES:
        chain = list(M.contracts_for(com, as_of))
        cons = contracts.get(com, {})
        chain = [_PRE[com] + cons[m] if m in cons else chain[i]
                 for i, m in enumerate(months)]
        cal[com] = list(zip(months, chain))
    cif2 = {c: {m: v for m, v in mv.items() if m in months}
            for c, mv in cif.items() if c in M.COMMODITIES}
    frt2 = {r: {m: v for m, v in mv.items() if m in months}
            for r, mv in freight.items() if r in M.FREIGHT_REGIONS}
    return cif2, frt2, cal


def _futures_complete(fut, front):
    """True only when EVERY commodity has a price for the front month — that's
    what the carry chart's net-of-interest leg anchors on."""
    return bool(fut) and all((fut.get(c) or {}).get(front) is not None
                             for c in M.COMMODITIES)


def _live_futures_spreads(cal, as_of, tries=3):
    """Live CBOT board from Massive + the spreads it implies, so the carry charts
    (gross-carry shape + the net-of-interest overlay) work for imported dates.

    The Bid Sheet's OWN futures are live Eikon formulas that arrive as #N/A in the
    emailed file, so they can't be read from the attachment — we pull the board
    from Massive instead (the same source the app uses for the live fed-funds
    rate). Retries a few times for a transient Massive blip. Returns ({}, {}) if
    Massive isn't configured; the completeness check is left to the caller."""
    try:
        import rf_massive_futures as MF
        if not MF.configured():
            return {}, {}
    except Exception:
        return {}, {}
    import time
    front = (cal.get(M.COMMODITIES[0]) or [(None,)])[0][0]
    fut = {}
    for attempt in range(max(1, tries)):
        try:
            fut = MF.futures_for_calendar(cal, as_of)
        except Exception:
            fut = {}
        if _futures_complete(fut, front):
            break
        if attempt < tries - 1:
            time.sleep(2)
    return fut, _spreads_for(cal, fut)


def _spreads_for(cal, fut):
    """Inter-contract spreads implied by a CBOT board, per commodity, in the
    calendar's own contract order. The carry charts need them stored."""
    spr = {}
    for c in M.COMMODITIES:
        cons = [ct for _m, ct in cal.get(c, [])]
        months = [m for m, _ in cal.get(c, [])]
        try:
            labels = M.spread_labels_for(c, cons)
            vals = M.spreads_from_futures(c, fut.get(c, {}), contracts=cons, months=months)
            spr[c] = list(zip(labels, vals))
        except Exception:
            pass
    return spr


# CBOT root per FOB commodity, and month letter -> calendar month, for naming a
# sheet's contracts (e.g. ZCZ26; Massive's ticker for it is ZCZ6).
_ROOT = {"Corn": "ZC", "Soybeans": "ZS", "Wheat": "ZW"}
_LETTER_MONTH = {"F": 1, "G": 2, "H": 3, "J": 4, "K": 5, "M": 6,
                 "N": 7, "Q": 8, "U": 9, "V": 10, "X": 11, "Z": 12}


def close_symbol(commodity, code, as_of):
    """The sheet's contract code ('CZ', 'SF', ...) -> the CBOT contract it meant on
    `as_of`: the first one with that month letter on or after the as-of month, the
    same front-with-that-letter rule the live Massive pull follows.
    ('Corn', 'CZ', 2026-10-09) -> 'ZCZ26'; ('Corn', 'CH', 2026-10-09) -> 'ZCH27'."""
    letter = str(code)[-1].upper()
    year = as_of.year if _LETTER_MONTH[letter] >= as_of.month else as_of.year + 1
    return f"{_ROOT[commodity]}{letter}{year % 100:02d}"


def futures_from_closes(cal, closes, as_of):
    """calendar + {contract: $/bu} -> {commodity: {month: $/bu}}. A contract with
    no close is left out, so the completeness check catches it. Pure."""
    out = {}
    for com, pairs in (cal or {}).items():
        if com not in _ROOT:
            continue
        row = {}
        for month, code in pairs:
            if code:
                price = closes.get(close_symbol(com, code, as_of))
                if price is not None:
                    row[month] = price
        if row:
            out[com] = row
    return out


def _closing_board(as_of, symbols):
    """{contract: $/bu} at `as_of`'s own settlement, from Massive's daily session
    bars (/aggs/<ticker>, resolution 1session, matched on session_end_date). The
    live snapshot can't give a past day.

    Not the basis tracker's FUTURES_PRICES. Until 2026-10-12 its daily captures
    for date D held the settlement of the trading day BEFORE D; reading it put
    10/06, 10/07 and 10/09 a day off, until they were re-saved from these bars.
    The tracker fixed its capture and rebuilt those rows that night, so the table
    is right by date now, but these bars don't depend on its 3:45 PM run, and its
    row for today can be the session's close until its next run. Days saved from
    a live Massive pull agree with these bars to ~1¢.

    Returns {} for anything Massive can't supply; the completeness check then
    refuses the save, and the next run retries."""
    try:
        import os
        import rf_massive_api as MA
        key = os.environ.get("MASSIVE_API_KEY", "").strip()
        if not key or not symbols:
            return {}
        tickers = {s: s[:3] + s[-1] for s in symbols}          # ZCZ26 -> ZCZ6
        bars = MA.get_daily_bars_many(sorted(set(tickers.values())), key)
    except Exception:
        return {}
    out = {}
    for sym, ticker in tickers.items():
        df = bars.get(ticker)
        if df is not None and as_of in df.index:
            out[sym] = round(float(df.loc[as_of, "settle"]) / 100.0, 4)
    return out


class FuturesUnavailable(RuntimeError):
    """Raised when the CBOT board couldn't be had (Massive down, or no close
    captured for a past day), so committing would save a carry-less day. The
    caller can retry later or override."""


def save_bidsheet(path, as_of=None, commit=False, require_futures=True, futures="live"):
    """Parse the Bid Sheet at `path` and (if commit) upsert it into the archive.
    Pass `as_of` (from the attachment filename / email date) — A1 is unreliable.
    Also stores the CBOT board + spreads so the carry charts work.

    `futures`: "live" (default) pulls the current board from Massive, which is
    right for today's sheet between the 1:20 PM close and the 7 PM reopen.
    "close" takes `as_of`'s own settlement from Massive's daily session bars
    instead: the right source for a PAST day, or for today after the reopen. The
    import uses it when it catches up a missed day or runs late.

    `require_futures` (default True): refuse to commit unless every commodity has
    a front-month price, so a Massive blip can't silently upload a day with no
    carry/net data. Pass False to force a CIF+freight-only save. Returns a
    summary dict (`futures_complete` says whether the board is whole).
    Idempotent: save_snapshot replaces the date's rows."""
    as_of, cif, freight, contracts = parse_bidsheet(path, as_of=as_of)
    cif2, frt2, cal = build_payload(as_of, cif, freight, contracts)
    if futures == "close":
        symbols = {close_symbol(com, code, as_of)
                   for com, pairs in cal.items() if com in _ROOT
                   for _m, code in pairs if code}
        fut = futures_from_closes(cal, _closing_board(as_of, symbols), as_of)
        spr = _spreads_for(cal, fut)
    else:
        fut, spr = _live_futures_spreads(cal, as_of)
    months = M.months_for(as_of)
    complete = _futures_complete(fut, months[0] if months else None)
    was_present = as_of.isoformat() in {str(d)[:10] for d in db.list_dates()}
    out = dict(as_of=as_of.isoformat(), was_present=was_present,
               n_cif=sum(len(v) for v in cif2.values()),
               n_frt=sum(len(v) for v in frt2.values()),
               n_fut=sum(len(v) for v in fut.values()),
               futures_complete=complete, futures_source=futures)
    if commit:
        if db._backend() == "sqlite":
            raise RuntimeError("Refusing to write to the SQLite fallback — set "
                               "USE_SNOWFLAKE + SNOWFLAKE_* (or DATABASE_URL).")
        if require_futures and not complete:
            source = ("Massive's live board" if futures != "close"
                      else "Massive's daily settlements")
            raise FuturesUnavailable(
                f"CBOT board did not come from {source} for {as_of} "
                f"(commodities with futures: {sorted(c for c in fut if fut[c])}). "
                "Refusing to upload a carry-less day; the next run retries, "
                "or save with require_futures=False.")
        db.save_snapshot(as_of.isoformat(), cif2, frt2, cal, futures=fut, spreads=spr)
        out["committed"] = True
    return out
