"""
hl_model.py
-----------
JPSI Futures High / Low Forecast Model — self-contained module.
Renders as a top-level tab inside the Grain Seasonal Dashboard.

Data sources:
  - data/High Low Model.xlsx  : fundamental analog data (production, usage, H/L history)
  - legacy_history.py         : pre-2022 settlement prices (corn/soy to 2008)
  - massive_api.py            : Massive REST API for 2022-present daily bars
"""

from __future__ import annotations

from datetime import date, datetime
from pathlib import Path

import numpy as np
import pandas as pd
import plotly.graph_objects as go
import streamlit as st

from legacy_history import get_legacy_series
from massive_api import get_daily_bars_many

# ── Paths ────────────────────────────────────────────────────────────────────
_HERE = Path(__file__).parent
HL_DATA_PATH = _HERE / "data" / "High Low Model.xlsx"

# ── Colours (kept from original H/L model for chart consistency) ─────────────
_BG      = "#16201a"
_SURF    = "#1a2620"
_BORDER  = "#253328"
_TEXT    = "#e8ede9"
_MUTED   = "#7a9485"
_GREEN   = "#5e7164"
_LT      = "#8db89a"
COL_HIGH = "#8db89a"
COL_LOW  = "#6fa8c4"
COL_GOLD = "#c4b456"
COL_PURP = "#9b89c4"
COL_RED  = "#e07070"

# ── Session-state key prefix (avoids collisions with grain dashboard keys) ────
_P = "hl_"   # e.g. hl_cz_jan1

# ── Inline CSS scoped to H/L card elements only ──────────────────────────────
_CSS = f"""
<style>
.hl-row-label{{font-size:0.72rem;text-transform:uppercase;letter-spacing:0.08em;
  color:{_MUTED};margin-bottom:6px;padding-top:4px;}}
.hl-card{{background:{_SURF};border:1px solid {_BORDER};border-radius:12px;
  padding:16px 18px 14px;text-align:center;height:100%;}}
.hl-card-high{{border-top:4px solid {COL_HIGH};}}
.hl-card-low{{border-top:4px solid {COL_LOW};}}
.hl-contract{{color:{_MUTED};font-size:0.65rem;text-transform:uppercase;
  letter-spacing:0.10em;margin-bottom:2px;}}
.hl-label{{color:{_MUTED};font-size:0.60rem;text-transform:uppercase;
  letter-spacing:0.08em;margin-bottom:8px;}}
.hl-price-high{{color:{COL_HIGH};font-size:2rem;font-weight:700;line-height:1;}}
.hl-price-low{{color:{COL_LOW};font-size:2rem;font-weight:700;line-height:1;}}
.hl-sub{{color:{_MUTED};font-size:0.76rem;margin-top:5px;}}
.hl-sec-hdr{{font-size:1.0rem;font-weight:600;color:{_TEXT};
  border-left:4px solid {_LT};padding-left:10px;margin:1.4rem 0 0.7rem 0;}}
.hl-sec-div{{border-bottom:1px solid {_BORDER};margin:6px 0 12px 0;
  color:{_MUTED};font-size:0.66rem;text-transform:uppercase;letter-spacing:0.10em;
  padding-bottom:4px;}}
.hl-ratio-badge{{background:{_SURF};border:1px solid {_BORDER};border-radius:6px;
  padding:3px 10px;display:inline-block;font-size:0.82rem;color:{_TEXT};margin-top:4px;}}
.hl-indicated-box{{background:{_SURF};border:1px solid {_BORDER};border-radius:8px;
  padding:10px 14px;margin-bottom:10px;font-size:0.88rem;}}
.hl-input-card{{background:{_SURF};border:1px solid {_BORDER};border-radius:12px;
  padding:20px 22px 18px;margin-bottom:4px;}}
.hl-input-card-title{{font-size:1.0rem;font-weight:700;color:{_TEXT};
  margin-bottom:14px;padding-bottom:10px;border-bottom:1px solid {_BORDER};}}
</style>
"""

_SECTION_EMOJI = {
    "Prod ≥ Use":         "🟢",
    "Prod < Use":         "🔴",
    "Prod ≥ Prev Yr Use": "🟢",
    "Prod < Prev Yr Use": "🔴",
    "No Crop Scare":      "🟢",
    "SA Crop Problem":    "🟡",
    "Crop Scare":         "🔴",
}
_SECTION_COLORS_LOW = {
    "Prod ≥ Use": COL_LOW, "Prod < Use": COL_GOLD,
    "Prod ≥ Prev Yr Use": COL_LOW, "Prod < Prev Yr Use": COL_GOLD,
    "No Crop Scare": COL_LOW, "SA Crop Problem": COL_GOLD,
    "Crop Scare": COL_RED, "All": COL_LOW,
}
_SECTION_COLORS_HIGH = {
    "Prod ≥ Use": COL_HIGH, "Prod < Use": COL_PURP,
    "Prod ≥ Prev Yr Use": COL_HIGH, "Prod < Prev Yr Use": COL_PURP,
    "No Crop Scare": COL_HIGH, "SA Crop Problem": COL_PURP,
    "Crop Scare": COL_RED, "All": COL_HIGH,
}


# ─────────────────────────────────────────────────────────────────────────────
# Data loading — High Low Model.xlsx
# ─────────────────────────────────────────────────────────────────────────────

def _parse_year(v):
    if v is None:
        return None
    try:
        return int(str(v).strip())
    except Exception:
        return None


def _is_num(v):
    return isinstance(v, (int, float)) and not isinstance(v, bool)


def _load_cz(wb, name):
    rows = list(wb[name].iter_rows(values_only=True))
    out = []
    for row in rows:
        yr = _parse_year(row[1])
        if not (yr and 1985 <= yr <= 2035):
            continue
        usage, prod, jan1, price = row[2], row[4], row[8], row[12]
        if not all(_is_num(v) for v in [usage, prod, jan1, price]):
            continue
        out.append(dict(
            year=yr, usage=usage, production=prod,
            ratio=prod / usage, jan1=jan1,
            date=row[10] if isinstance(row[10], datetime) else None,
            price=price, price_pct=price / jan1,
            is_est=(row[13] == "Est"),
        ))
    df = pd.DataFrame(out)
    if not df.empty:
        df["section"] = df["ratio"].apply(lambda r: "Prod ≥ Use" if r >= 1 else "Prod < Use")
    return df


def _load_cn(wb, name):
    rows = list(wb[name].iter_rows(values_only=True))
    out = []
    for i, row in enumerate(rows):
        if i >= 48:
            break
        yr = _parse_year(row[1])
        if not (yr and 1985 <= yr <= 2035):
            continue
        co_pct, price, jan1 = row[4], row[5], row[9]
        if not all(_is_num(v) for v in [co_pct, price, jan1]):
            continue
        out.append(dict(
            year=yr, carryout_pct=co_pct, jan1=jan1,
            date=row[7] if isinstance(row[7], datetime) else None,
            price=price, price_pct=price / jan1,
            is_est=(row[6] == "Est"),
        ))
    return pd.DataFrame(out)


def _load_sx(wb, name):
    rows = list(wb[name].iter_rows(values_only=True))
    out = []
    section = "Prod ≥ Prev Yr Use"
    for i, row in enumerate(rows):
        if i < 8:
            continue
        rstr = " ".join(str(v) for v in row if v is not None)
        if "Less Than" in rstr:
            section = "Prod < Prev Yr Use"
            continue
        if "Equal to or Greater" in rstr:
            section = "Prod ≥ Prev Yr Use"
            continue
        yr = _parse_year(row[0])
        if not (yr and 1985 <= yr <= 2035):
            continue
        co_pct, jan1, price = row[1], row[3], row[7]
        if not all(_is_num(v) for v in [co_pct, jan1, price]):
            continue
        out.append(dict(
            year=yr, carryout_pct=co_pct, jan1=jan1,
            date=row[5] if isinstance(row[5], datetime) else None,
            price=price, price_pct=price / jan1,
            is_est=(row[8] == "Est"),
            section=section,
        ))
    return pd.DataFrame(out)


def _load_sn(wb, name, is_low):
    rows = list(wb[name].iter_rows(values_only=True))
    out = []
    section = "No Crop Scare"
    yr_idx = 1 if is_low else 0
    co_idx = 6 if is_low else 5
    for i, row in enumerate(rows):
        if i < 9:
            continue
        rstr = " ".join(str(v) for v in row if v is not None).lower()
        if "without" in rstr and "crop scare" in rstr:
            section = "No Crop Scare"; continue
        if "south america" in rstr:
            section = "SA Crop Problem"; continue
        if "with" in rstr and "crop scare" in rstr and "without" not in rstr:
            section = "Crop Scare"; continue
        yr = _parse_year(row[yr_idx])
        if not (yr and 1985 <= yr <= 2035):
            continue
        co_pct, price, jan1 = row[co_idx], row[7], row[11]
        if not all(_is_num(v) for v in [co_pct, price, jan1]):
            continue
        out.append(dict(
            year=yr, carryout_pct=co_pct, jan1=jan1,
            date=row[9] if isinstance(row[9], datetime) else None,
            price=price, price_pct=price / jan1,
            is_est=(row[8] == "Est"),
            section=section,
        ))
    return pd.DataFrame(out)


@st.cache_data(show_spinner=False)
def load_hl_data() -> dict:
    import openpyxl
    wb = openpyxl.load_workbook(HL_DATA_PATH)
    return {
        "cz_low":  _load_cz(wb, "CZ  Low"),
        "cz_high": _load_cz(wb, "CZ High"),
        "cn_low":  _load_cn(wb, "CN Low"),
        "cn_high": _load_cn(wb, "CN High"),
        "sx_low":  _load_sx(wb, "SX Low"),
        "sx_high": _load_sx(wb, "SX High"),
        "sn_low":  _load_sn(wb, "SN Low",  is_low=True),
        "sn_high": _load_sn(wb, "SN High", is_low=False),
    }


# ─────────────────────────────────────────────────────────────────────────────
# Seasonal price history — builds {year: pd.Series(close, index=date)}
# from legacy CSV (2008-2021) + Massive REST API (2022-present)
# ─────────────────────────────────────────────────────────────────────────────

@st.cache_data(ttl="6h", show_spinner=False)
def _load_hl_bars(tickers: tuple[str, ...], api_key: str) -> dict[str, pd.DataFrame]:
    return get_daily_bars_many(list(tickers), api_key)


def _build_seasonal_contracts(
    product_code: str,
    month_letter: str,
    api_key: str,
    legacy_df: pd.DataFrame,
    current_year: int,
    fallback_letter: str | None = None,
    fallback_end_month: int | None = None,
) -> dict[int, pd.Series]:
    """
    Return {year: Series(close_price, index=date)} for use in build_seasonal_df_hl().
    Merges legacy history (corn/soy 2008-2021) with Massive bars (2022-present).
    For contracts that lack their own history (e.g. Jul corn = ZCN), falls back to
    the Dec contract (ZCZ) filtered to fallback_end_month.
    """
    contracts: dict[int, pd.Series] = {}

    # ── Legacy years 2008-2021 ────────────────────────────────────────────────
    for year in range(2008, 2022):
        series = get_legacy_series(legacy_df, product_code, month_letter, year)
        if series is not None and len(series) > 5:
            contracts[year] = series
        elif fallback_letter:
            fb = get_legacy_series(legacy_df, product_code, fallback_letter, year)
            if fb is not None and len(fb) > 5:
                if fallback_end_month:
                    fb = fb[pd.DatetimeIndex(fb.index).month <= fallback_end_month]
                if len(fb) > 5:
                    contracts[year] = fb

    # ── Recent years 2022-present via Massive ─────────────────────────────────
    recent_tickers = []
    for year in range(2022, current_year + 1):
        digit = year % 10
        recent_tickers.append(f"{product_code}{month_letter}{digit}")
        if fallback_letter:
            recent_tickers.append(f"{product_code}{fallback_letter}{digit}")

    bars = _load_hl_bars(tuple(recent_tickers), api_key)

    for year in range(2022, current_year + 1):
        digit = year % 10
        primary = f"{product_code}{month_letter}{digit}"
        df = bars.get(primary)
        if df is not None and not df.empty:
            s = df["settle"]
            s = s[pd.DatetimeIndex(s.index).year == year]
            if len(s) > 5:
                contracts[year] = s
                continue
        if fallback_letter:
            fb_ticker = f"{product_code}{fallback_letter}{digit}"
            fb_df = bars.get(fb_ticker)
            if fb_df is not None and not fb_df.empty:
                s = fb_df["settle"]
                s = s[pd.DatetimeIndex(s.index).year == year]
                if fallback_end_month:
                    s = s[pd.DatetimeIndex(s.index).month <= fallback_end_month]
                if len(s) > 5:
                    contracts[year] = s

    return contracts


def build_seasonal_df_hl(
    contracts: dict[int, pd.Series],
    end_month: int = 12,
) -> pd.DataFrame:
    """Long-format seasonal DataFrame from {year: Series(close)}."""
    rows = []
    for year, series in contracts.items():
        s = series[pd.DatetimeIndex(series.index).month <= end_month]
        if len(s) < 10:
            continue
        jan1 = float(s.iloc[0])
        if pd.isna(jan1) or jan1 <= 0:
            continue
        for dt, close in s.items():
            if pd.isna(close):
                continue
            dt = pd.Timestamp(dt)
            rows.append({
                "year":      year,
                "date":      dt,
                "doy":       dt.timetuple().tm_yday,
                "close_raw": float(close),
                "price_pct": float(close) / jan1 * 100,
            })
    return pd.DataFrame(rows)


# ─────────────────────────────────────────────────────────────────────────────
# Computation helpers
# ─────────────────────────────────────────────────────────────────────────────

def _rank_df(df, ratio_col, current_ratio, current_jan1):
    d = df[~df["is_est"]].copy()
    d["distance"] = (d[ratio_col] - current_ratio).abs()
    d = d.sort_values("distance").reset_index(drop=True)
    d.index = d.index + 1
    d["indicated"] = (d["price_pct"] * current_jan1).round(2)
    return d


def _headline(df, ratio_col, current_ratio, current_jan1):
    d = _rank_df(df, ratio_col, current_ratio, current_jan1)
    if d.empty:
        return None, None, None
    med_pct  = d["price_pct"].median()
    top5_pct = d.head(5)["price_pct"].median()
    return round(med_pct * current_jan1, 2), med_pct, round(top5_pct * current_jan1, 2)


def _fmt_pct(p):
    return f"{p*100:.1f}%" if p is not None else "—"


def _fmt_price(p):
    return f"${p/100:.2f}/bu" if p is not None else "—"


def _build_display_table(df, ratio_col, ratio_label, current_ratio, current_jan1,
                          price_col_label, current_section=None):
    d = _rank_df(df, ratio_col, current_ratio, current_jan1)
    if d.empty:
        return pd.DataFrame()
    out = pd.DataFrame()
    out["Rank"] = d.index
    out["Year"] = d["year"].astype(int)
    if "section" in d.columns:
        out["Category"] = d["section"].apply(lambda s: f"{_SECTION_EMOJI.get(s,'⚪')} {s}")
        if current_section:
            out["Cur Yr"] = d["section"].apply(lambda s: "★" if s == current_section else "")
    out[ratio_label]          = (d[ratio_col] * 100).round(2).astype(str) + "%"
    out["Dist. from Current"] = (d["distance"] * 100).round(3).astype(str) + "%"
    out["Jan 1 (¢/bu)"]       = d["jan1"].round(2)
    out[price_col_label]       = d["price"].round(2)
    out["% of Jan 1"]          = (d["price_pct"] * 100).round(1).astype(str) + "%"
    out["Indicated ($/bu)"]    = (d["indicated"] / 100).round(4)
    out["Date of H/L"]         = d["date"].apply(
        lambda x: x.strftime("%m/%d/%Y") if (pd.notna(x) and isinstance(x, datetime)) else "—"
    )
    return out


# ─────────────────────────────────────────────────────────────────────────────
# Chart builders
# ─────────────────────────────────────────────────────────────────────────────

def _layout_base(title="", height=430):
    return dict(
        title=dict(text=title, font=dict(color=_TEXT, size=12), x=0),
        plot_bgcolor=_SURF, paper_bgcolor=_SURF,
        font=dict(color=_TEXT, family="Arial"),
        height=height,
        margin=dict(l=60, r=20, t=40, b=50),
        legend=dict(bgcolor=_BG, bordercolor=_BORDER, borderwidth=1,
                    font=dict(color=_TEXT, size=10)),
    )


def _make_scatter(df_low, df_high, ratio_col, ratio_label, current_ratio, title):
    fig = go.Figure()
    for df, lbl, cmap, sym in [
        (df_low,  "Low",  _SECTION_COLORS_LOW,  "circle"),
        (df_high, "High", _SECTION_COLORS_HIGH, "diamond"),
    ]:
        d = df[~df["is_est"]].copy()
        groups = [(s, d[d["section"] == s]) for s in d["section"].unique()] \
            if "section" in d.columns else [("All", d)]
        for sec, sd in groups:
            if sd.empty:
                continue
            fig.add_trace(go.Scatter(
                x=sd[ratio_col] * 100, y=sd["price_pct"] * 100,
                mode="markers+text",
                text=sd["year"].astype(str),
                textposition="top center",
                textfont=dict(size=8, color=_MUTED),
                marker=dict(size=9, color=cmap.get(sec, _MUTED), symbol=sym,
                            line=dict(width=1, color=_BORDER)),
                name=f"{lbl} — {sec}",
                hovertemplate=f"<b>%{{text}}</b> ({lbl})<br>{ratio_label}: %{{x:.2f}}%<br>"
                              "Price % of Jan 1: %{y:.1f}%<extra></extra>",
            ))
    fig.add_vline(
        x=current_ratio * 100, line_color=_LT, line_dash="dash", line_width=2,
        annotation_text=f" Current: {current_ratio*100:.2f}%",
        annotation_position="top right",
        annotation_font=dict(color=_LT, size=11),
    )
    fig.update_layout(
        **_layout_base(title),
        xaxis=dict(title=ratio_label, gridcolor=_BORDER, color=_MUTED, ticksuffix="%", zeroline=False),
        yaxis=dict(title="Price as % of Jan 1", gridcolor=_BORDER, color=_MUTED, ticksuffix="%", zeroline=False),
    )
    return fig


_MONTH_LABELS = ["Jan","Feb","Mar","Apr","May","Jun","Jul","Aug","Sep","Oct","Nov","Dec"]

def _make_seasonality_chart(df_low, df_high, title):
    def month_counts(df):
        dates = df[~df["is_est"]]["date"]
        months = dates.apply(lambda x: x.month if (pd.notna(x) and isinstance(x, datetime)) else None).dropna()
        return [int((months == m).sum()) for m in range(1, 13)]
    fig = go.Figure()
    fig.add_trace(go.Bar(x=_MONTH_LABELS, y=month_counts(df_low),  name="Annual Low",  marker_color=COL_LOW,  marker_line=dict(width=0)))
    fig.add_trace(go.Bar(x=_MONTH_LABELS, y=month_counts(df_high), name="Annual High", marker_color=COL_HIGH, marker_line=dict(width=0)))
    fig.update_layout(
        **_layout_base(title, height=320),
        barmode="group",
        xaxis=dict(title="Month", gridcolor=_BORDER, color=_MUTED, zeroline=False),
        yaxis=dict(title="Number of Years", gridcolor=_BORDER, color=_MUTED, zeroline=False, dtick=1),
    )
    return fig


_MONTH_START_DOY  = [1, 32, 60, 91, 121, 152, 182, 213, 244, 274, 305, 335]
_MONTH_END_DOY    = [31, 59, 90, 120, 151, 181, 212, 243, 273, 304, 334, 365]
_ALL_MONTH_NAMES  = ["Jan","Feb","Mar","Apr","May","Jun","Jul","Aug","Sep","Oct","Nov","Dec"]


def _make_seasonal_overlay(
    seasonal_df, current_year, title,
    end_month=12,
    ind_high_price=None, ind_low_price=None,
    jan1_price=None,
    show_individual_years=True,
    height=500,
):
    """
    MRCI-style seasonal chart. jan1_price in ¢/bu; ind_high/low_price in $/bu.
    All y-axis values rendered in $/bu.
    """
    j = (jan1_price or 100) / 100 / 100   # ¢/bu → $/bu factor

    def pct_to_price(pct_series):
        return pct_series * j

    fig = go.Figure()

    hist = seasonal_df[seasonal_df["year"] != current_year].copy()
    curr = seasonal_df[seasonal_df["year"] == current_year].copy()

    # Month boundary vlines
    for m_idx in range(1, end_month):
        fig.add_vline(x=_MONTH_START_DOY[m_idx],
                      line_color="rgba(160,140,200,0.20)", line_width=1, line_dash="dot")

    # Faint historical year lines
    if show_individual_years:
        for yr, g in hist.groupby("year"):
            g = g.sort_values("doy")
            fig.add_trace(go.Scatter(
                x=g["doy"], y=pct_to_price(g["price_pct"]),
                mode="lines", line=dict(width=0.6, color="rgba(120,140,130,0.22)"),
                name=str(yr), showlegend=False,
                hovertemplate=f"<b>{yr}</b>  $%{{y:.2f}}/bu<extra></extra>",
            ))

    def _smooth(s):
        return s.rolling(window=7, center=True, min_periods=1).mean()

    def _avg_price(df):
        return _smooth(df.groupby("doy")["price_pct"].mean()).apply(lambda v: v * j)

    def _std_price(df):
        return _smooth(df.groupby("doy")["price_pct"].std().fillna(0)).apply(lambda v: v * j)

    avg_full = _avg_price(hist)
    std_full = _std_price(hist)
    doys     = avg_full.index.tolist()
    n_full   = hist["year"].nunique()
    lw       = 2.0 if show_individual_years else 2.4

    # Full avg ± 1 SD band
    fig.add_trace(go.Scatter(
        x=doys + doys[::-1],
        y=list(avg_full + std_full) + list((avg_full - std_full).iloc[::-1]),
        fill="toself", fillcolor="rgba(94,113,100,0.13)",
        line=dict(width=0), name="±1 SD", hoverinfo="skip", showlegend=True,
    ))
    fig.add_trace(go.Scatter(
        x=doys, y=avg_full.values, mode="lines",
        line=dict(width=lw, color=_LT),
        name=f"Full Avg ({n_full} yr)",
        hovertemplate="Full avg: $%{y:.2f}/bu<extra></extra>",
    ))

    # 5-yr avg
    hist5 = hist[hist["year"] >= current_year - 5]
    if hist5["year"].nunique() >= 2:
        avg5 = _avg_price(hist5)
        fig.add_trace(go.Scatter(
            x=avg5.index.tolist(), y=avg5.values, mode="lines",
            line=dict(width=lw, color=COL_RED, dash="dash"),
            name=f"5-Yr Avg ({hist5['year'].nunique()} yr)",
            hovertemplate="5-yr avg: $%{y:.2f}/bu<extra></extra>",
        ))

    # 15-yr avg
    hist15 = hist[hist["year"] >= current_year - 15]
    if hist15["year"].nunique() >= 3:
        avg15 = _avg_price(hist15)
        fig.add_trace(go.Scatter(
            x=avg15.index.tolist(), y=avg15.values, mode="lines",
            line=dict(width=lw, color=COL_LOW, dash="dot"),
            name=f"15-Yr Avg ({hist15['year'].nunique()} yr)",
            hovertemplate="15-yr avg: $%{y:.2f}/bu<extra></extra>",
        ))

    # Current year — actual prices
    if not curr.empty:
        curr = curr.sort_values("doy")
        fig.add_trace(go.Scatter(
            x=curr["doy"], y=curr["close_raw"] / 100,
            mode="lines", line=dict(width=3.0 if not show_individual_years else 2.8, color=COL_GOLD),
            name=f"{current_year}  (Current)",
            hovertemplate=f"<b>{current_year}</b>  $%{{y:.2f}}/bu<extra></extra>",
        ))

    # Tight y-range
    rescaled    = pct_to_price(hist["price_pct"].dropna())
    curr_prices = (curr["close_raw"] / 100) if not curr.empty else pd.Series([], dtype=float)
    candidate_y = pd.concat([rescaled, curr_prices, avg_full.dropna()], ignore_index=True).dropna()

    if len(candidate_y) > 0:
        y_lo = float(np.percentile(candidate_y, 2))
        y_hi = float(np.percentile(candidate_y, 98))
        if not curr_prices.empty:
            y_lo = min(y_lo, float(curr_prices.min()))
            y_hi = max(y_hi, float(curr_prices.max()))
        if ind_high_price is not None:
            y_hi = max(y_hi, ind_high_price)
        if ind_low_price is not None:
            y_lo = min(y_lo, ind_low_price)
        jan1_d = (jan1_price or 100) / 100
        y_lo = min(y_lo, jan1_d)
        y_hi = max(y_hi, jan1_d)
        pad = (y_hi - y_lo) * 0.09
        y_range = [y_lo - pad, y_hi + pad]
    else:
        jan1_d = (jan1_price or 100) / 100
        y_range = [jan1_d * 0.85, jan1_d * 1.20]

    # Today band
    today_doy = datetime.now().timetuple().tm_yday
    fig.add_vrect(x0=today_doy - 1.5, x1=today_doy + 1.5,
                  fillcolor="rgba(80,200,220,0.22)", line_width=0,
                  annotation_text="Today", annotation_position="top",
                  annotation_font=dict(color="rgba(80,200,220,0.80)", size=9))

    # Jan 1 baseline
    jan1_d = (jan1_price or 100) / 100
    fig.add_hline(y=jan1_d, line_dash="dot", line_color=_MUTED, line_width=1,
                  annotation_text=f" Jan 1  ${jan1_d:.2f}",
                  annotation_position="right",
                  annotation_font=dict(color=_MUTED, size=9))

    # Indicated H/L lines
    if ind_high_price is not None:
        fig.add_hline(y=ind_high_price, line_dash="dash", line_color=COL_HIGH, line_width=1.8,
                      annotation_text=f" Ind. High  ${ind_high_price:.2f}",
                      annotation_position="right",
                      annotation_font=dict(color=COL_HIGH, size=10))
    if ind_low_price is not None:
        fig.add_hline(y=ind_low_price, line_dash="dash", line_color=COL_LOW, line_width=1.8,
                      annotation_text=f" Ind. Low  ${ind_low_price:.2f}",
                      annotation_position="right",
                      annotation_font=dict(color=COL_LOW, size=10))

    month_doys  = [d for d, m in zip(_MONTH_START_DOY, range(1, 13)) if m <= end_month]
    month_names = [n for n, m in zip(_ALL_MONTH_NAMES, range(1, 13)) if m <= end_month]
    x_max = _MONTH_END_DOY[end_month - 1] + 5

    layout = _layout_base(title, height=height)
    layout["margin"] = dict(l=70, r=140, t=45, b=50)  # wider right margin for the H/L/Jan-1 annotations
    fig.update_layout(
        **layout,
        xaxis=dict(title="Month", tickvals=month_doys, ticktext=month_names,
                   gridcolor=_BORDER, color=_MUTED, zeroline=False, range=[1, x_max]),
        yaxis=dict(title="Price ($/bu)", gridcolor=_BORDER, color=_MUTED,
                   zeroline=False, tickprefix="$", range=y_range, autorange=False),
    )
    return fig


# ─────────────────────────────────────────────────────────────────────────────
# HTML tile helpers
# ─────────────────────────────────────────────────────────────────────────────

def _headline_tile(contract, label, price, pct, top5, kind):
    cls        = "high" if kind == "high" else "low"
    price_html = _fmt_price(price)
    pct_html   = f"Median: {_fmt_pct(pct)} of Jan 1" if pct else "—"
    top5_html  = f"Top 5 Comps: {_fmt_price(top5)}" if top5 else ""
    return f"""
<div class="hl-card hl-card-{cls}">
  <div class="hl-contract">{contract}</div>
  <div class="hl-label">{label}</div>
  <div class="hl-price-{cls}">{price_html}</div>
  <div class="hl-sub">{pct_html}</div>
  <div class="hl-sub">{top5_html}</div>
</div>"""


# ─────────────────────────────────────────────────────────────────────────────
# Per-contract sub-tab renderer
# ─────────────────────────────────────────────────────────────────────────────

def _render_contract_tab(
    label, D_low, D_high, ratio_col, ratio_label, current_ratio,
    current_jan1, hl_low, lp, t5_low, hl_high, hp, t5_high,
    price_low_label, price_high_label, scatter_title, season_title,
    current_section=None, tile_low_label="Indicated Low", tile_high_label="Indicated High",
):
    t1, t2 = st.columns(2)
    t1.markdown(_headline_tile(label, tile_low_label,  hl_low,  lp,  t5_low,  "low"),  unsafe_allow_html=True)
    t2.markdown(_headline_tile(label, tile_high_label, hl_high, hp, t5_high, "high"), unsafe_allow_html=True)
    st.markdown('<div style="height:12px;"></div>', unsafe_allow_html=True)

    st.markdown("#### Historical Scatter: Ratio vs Price as % of Jan 1")
    st.plotly_chart(_make_scatter(D_low, D_high, ratio_col, ratio_label, current_ratio, scatter_title),
                    use_container_width=True)

    st.markdown("#### Timing: Month of Annual High / Low")
    st.plotly_chart(_make_seasonality_chart(D_low, D_high, season_title), use_container_width=True)

    col_l, col_h = st.columns(2)
    with col_l:
        tbl = _build_display_table(D_low, ratio_col, ratio_label, current_ratio, current_jan1,
                                    price_low_label, current_section=current_section)
        st.markdown(
            f'<div class="hl-indicated-box">📉 <b>Lows</b> ranked by similarity &nbsp;|&nbsp; '
            f'Median: <b style="color:{COL_LOW}">{_fmt_price(hl_low)}</b> '
            f'({_fmt_pct(lp)} of Jan 1) &nbsp;|&nbsp; '
            f'Top-5: <b style="color:{COL_LOW}">{_fmt_price(t5_low)}</b>'
            + (f'&nbsp;·&nbsp; ★ = {current_section}' if current_section else '') + '</div>',
            unsafe_allow_html=True,
        )
        st.dataframe(tbl, use_container_width=True, hide_index=True, height=480)
    with col_h:
        tbl = _build_display_table(D_high, ratio_col, ratio_label, current_ratio, current_jan1,
                                    price_high_label, current_section=current_section)
        st.markdown(
            f'<div class="hl-indicated-box">📈 <b>Highs</b> ranked by similarity &nbsp;|&nbsp; '
            f'Median: <b style="color:{COL_HIGH}">{_fmt_price(hl_high)}</b> '
            f'({_fmt_pct(hp)} of Jan 1) &nbsp;|&nbsp; '
            f'Top-5: <b style="color:{COL_HIGH}">{_fmt_price(t5_high)}</b>'
            + (f'&nbsp;·&nbsp; ★ = {current_section}' if current_section else '') + '</div>',
            unsafe_allow_html=True,
        )
        st.dataframe(tbl, use_container_width=True, hide_index=True, height=480)


# ─────────────────────────────────────────────────────────────────────────────
# Main render entry point
# ─────────────────────────────────────────────────────────────────────────────

def render_hl_tab(api_key: str, legacy_df) -> None:
    """Render the full H/L Forecast Model inside a Streamlit tab."""
    st.markdown(_CSS, unsafe_allow_html=True)

    # ── Check data file ───────────────────────────────────────────────────────
    if not HL_DATA_PATH.exists():
        st.error(
            f"H/L Model data file not found: `{HL_DATA_PATH}`\n\n"
            "Copy `High Low Model.xlsx` into the `data/` folder."
        )
        return

    try:
        D = load_hl_data()
    except Exception as e:
        st.error(f"Error loading H/L model data: {e}")
        return

    # ── Read assumptions from session state (defaults applied on first run) ───
    # Keys are prefixed with "hl_" to avoid collisions with grain dashboard keys.
    cz_jan1    = st.session_state.get(f"{_P}cz_jan1",    458.50)
    cz_prod    = st.session_state.get(f"{_P}cz_prod",  15979.0)
    cz_use     = st.session_state.get(f"{_P}cz_use",   16120.0)
    cn_jan1    = st.session_state.get(f"{_P}cn_jan1",    452.00)
    cn_co_raw  = st.session_state.get(f"{_P}cn_co",       12.9)
    sx_jan1    = st.session_state.get(f"{_P}sx_jan1",  1062.75)
    sx_co_raw  = st.session_state.get(f"{_P}sx_co",       28.4)
    sx_section = st.session_state.get(f"{_P}sx_section", "Prod ≥ Prev Yr Use")
    sn_jan1    = st.session_state.get(f"{_P}sn_jan1",  1072.00)
    sn_co_raw  = st.session_state.get(f"{_P}sn_co",       29.5)
    sn_section = st.session_state.get(f"{_P}sn_section", "No Crop Scare")

    cz_ratio  = cz_prod / cz_use if cz_use else 0.0
    cn_co_pct = cn_co_raw / 100
    sx_co_pct = sx_co_raw / 100
    sn_co_pct = sn_co_raw / 100
    cz_section = "Prod ≥ Use" if cz_ratio >= 1 else "Prod < Use"

    # ── Headline computations ─────────────────────────────────────────────────
    cz_hl_low,  cz_lp,  cz_t5_low  = _headline(D["cz_low"],  "ratio",        cz_ratio,  cz_jan1)
    cz_hl_high, cz_hp,  cz_t5_high = _headline(D["cz_high"], "ratio",        cz_ratio,  cz_jan1)
    cn_hl_low,  cn_lp,  cn_t5_low  = _headline(D["cn_low"],  "carryout_pct", cn_co_pct, cn_jan1)
    cn_hl_high, cn_hp,  cn_t5_high = _headline(D["cn_high"], "carryout_pct", cn_co_pct, cn_jan1)
    sx_hl_low,  sx_lp,  sx_t5_low  = _headline(D["sx_low"],  "carryout_pct", sx_co_pct, sx_jan1)
    sx_hl_high, sx_hp,  sx_t5_high = _headline(D["sx_high"], "carryout_pct", sx_co_pct, sx_jan1)
    sn_hl_low,  sn_lp,  sn_t5_low  = _headline(D["sn_low"],  "carryout_pct", sn_co_pct, sn_jan1)
    sn_hl_high, sn_hp,  sn_t5_high = _headline(D["sn_high"], "carryout_pct", sn_co_pct, sn_jan1)

    def _to_dollar(cents):
        return round(cents / 100, 4) if cents else None

    # ── Inner tabs ────────────────────────────────────────────────────────────
    (tab_ov, tab_inp, tab_cz, tab_cn, tab_sx, tab_sn, tab_seas) = st.tabs([
        "📌 Overview",
        "⚙️ Assumptions",
        "🌽 Dec Corn (CZ)",
        "🌽 Jul Corn (CN)",
        "🫘 Nov Soybeans (SX)",
        "🫘 Jul Soybeans (SN)",
        "📈 Seasonals",
    ])

    # ══ OVERVIEW ══════════════════════════════════════════════════════════════
    with tab_ov:
        st.markdown('<div class="hl-sec-hdr">📌 Indicated Price Range — Current Marketing Year</div>',
                    unsafe_allow_html=True)

        st.markdown('<div class="hl-row-label">📉 Indicated Lows</div>', unsafe_allow_html=True)
        c1, c2, c3, c4 = st.columns(4)
        c1.markdown(_headline_tile("CZ — Dec Corn",  "Jan–Dec Indicated Low",  cz_hl_low,  cz_lp,  cz_t5_low,  "low"),  unsafe_allow_html=True)
        c2.markdown(_headline_tile("CN — Jul Corn",  "Jan–Jul Indicated Low",  cn_hl_low,  cn_lp,  cn_t5_low,  "low"),  unsafe_allow_html=True)
        c3.markdown(_headline_tile("SX — Nov Beans", "Jan–Nov Indicated Low",  sx_hl_low,  sx_lp,  sx_t5_low,  "low"),  unsafe_allow_html=True)
        c4.markdown(_headline_tile("SN — Jul Beans", "Jan–Jul Indicated Low",  sn_hl_low,  sn_lp,  sn_t5_low,  "low"),  unsafe_allow_html=True)

        st.markdown('<div style="height:8px;"></div>', unsafe_allow_html=True)

        st.markdown('<div class="hl-row-label">📈 Indicated Highs</div>', unsafe_allow_html=True)
        c1, c2, c3, c4 = st.columns(4)
        c1.markdown(_headline_tile("CZ — Dec Corn",  "Jan–Dec Indicated High", cz_hl_high, cz_hp, cz_t5_high, "high"), unsafe_allow_html=True)
        c2.markdown(_headline_tile("CN — Jul Corn",  "Jan–Jul Indicated High", cn_hl_high, cn_hp, cn_t5_high, "high"), unsafe_allow_html=True)
        c3.markdown(_headline_tile("SX — Nov Beans", "Jan–Nov Indicated High", sx_hl_high, sx_hp, sx_t5_high, "high"), unsafe_allow_html=True)
        c4.markdown(_headline_tile("SN — Jul Beans", "Jan–Jul Indicated High", sn_hl_high, sn_hp, sn_t5_high, "high"), unsafe_allow_html=True)

        st.markdown('<div style="height:16px;"></div>', unsafe_allow_html=True)
        st.markdown('<div class="hl-sec-hdr">📋 Current Assumptions Summary</div>', unsafe_allow_html=True)
        s1, s2, s3, s4 = st.columns(4)
        sn_color = {"No Crop Scare": _LT, "SA Crop Problem": COL_GOLD, "Crop Scare": COL_RED}.get(sn_section, _MUTED)
        sx_color = _LT if "≥" in sx_section else COL_RED
        for col, contract, j1, metric_label, metric_val, extra in [
            (s1, "🌽 Dec Corn (CZ)",  cz_jan1, "Prod/Use",      f"{cz_ratio*100:.2f}% ({cz_section})", None),
            (s2, "🌽 Jul Corn (CN)",  cn_jan1, "Carryout/Use",   f"{cn_co_pct*100:.2f}%",              None),
            (s3, "🫘 Nov Beans (SX)", sx_jan1, "World C/O/Use",  f"{sx_co_pct*100:.2f}% ({sx_section})", sx_color),
            (s4, "🫘 Jul Beans (SN)", sn_jan1, "World C/O/Use",  f"{sn_co_pct*100:.2f}% ({sn_section})", sn_color),
        ]:
            col.markdown(
                f'<div class="hl-card" style="text-align:left;">'
                f'<div class="hl-contract" style="margin-bottom:8px;">{contract}</div>'
                f'<div style="font-size:0.82rem;color:{_MUTED};">Jan 1 Price</div>'
                f'<div style="font-size:1.1rem;font-weight:600;color:{_TEXT};">${j1/100:.2f}/bu</div>'
                f'<div style="font-size:0.82rem;color:{_MUTED};margin-top:6px;">{metric_label}</div>'
                f'<div style="font-size:1.0rem;color:{_TEXT};">{metric_val}</div>'
                f'</div>', unsafe_allow_html=True,
            )

    # ══ ASSUMPTIONS ═══════════════════════════════════════════════════════════
    with tab_inp:
        st.markdown(
            f'<div class="hl-sec-hdr">⚙️ Current Year Assumptions</div>'
            f'<div style="color:{_MUTED};font-size:0.82rem;margin-bottom:20px;">'
            f'Update these values to recalculate all indicated highs and lows.</div>',
            unsafe_allow_html=True,
        )
        col_a, col_b = st.columns(2)
        with col_a:
            st.markdown('<div class="hl-input-card"><div class="hl-input-card-title">🌽 Dec Corn (CZ26)</div></div>', unsafe_allow_html=True)
            cz_j1 = st.number_input("Jan 1 Price (¢/bu)", value=float(cz_jan1), step=0.25, key=f"{_P}cz_jan1", format="%.2f")
            cz_pr = st.number_input("US Production (Mil Bu)", value=float(cz_prod), step=50.0, key=f"{_P}cz_prod", format="%.0f")
            cz_us = st.number_input("US Usage (Mil Bu)",      value=float(cz_use),  step=50.0, key=f"{_P}cz_use",  format="%.0f")
            _r = cz_pr / cz_us if cz_us else 0.0
            st.markdown(f'<span class="hl-ratio-badge">Prod/Use: {_r*100:.2f}%  ·  {"Prod ≥ Use" if _r>=1 else "Prod < Use"}</span>', unsafe_allow_html=True)
        with col_b:
            st.markdown('<div class="hl-input-card"><div class="hl-input-card-title">🌽 Jul Corn (CN26)</div></div>', unsafe_allow_html=True)
            st.number_input("Jan 1 Price (¢/bu)", value=float(cn_jan1), step=0.25, key=f"{_P}cn_jan1", format="%.2f")
            cn_co = st.number_input("Carryout % of Use", value=float(cn_co_raw), step=0.1, key=f"{_P}cn_co", format="%.1f", help="e.g. 12.9 for 12.9%")
            st.markdown(f'<span class="hl-ratio-badge">C/O Ratio: {cn_co:.2f}%</span>', unsafe_allow_html=True)

        st.markdown('<div style="height:18px;"></div>', unsafe_allow_html=True)
        col_c, col_d = st.columns(2)
        with col_c:
            st.markdown('<div class="hl-input-card"><div class="hl-input-card-title">🫘 Nov Soybeans (SX26)</div></div>', unsafe_allow_html=True)
            st.number_input("Jan 1 Price (¢/bu)", value=float(sx_jan1), step=0.25, key=f"{_P}sx_jan1", format="%.2f")
            sx_co = st.number_input("World C/O % of Use", value=float(sx_co_raw), step=0.1, key=f"{_P}sx_co", format="%.1f")
            st.selectbox("Current Year Category", ["Prod ≥ Prev Yr Use", "Prod < Prev Yr Use"], key=f"{_P}sx_section")
            st.markdown(f'<span class="hl-ratio-badge">World C/O: {sx_co:.2f}%</span>', unsafe_allow_html=True)
        with col_d:
            st.markdown('<div class="hl-input-card"><div class="hl-input-card-title">🫘 Jul Soybeans (SN26)</div></div>', unsafe_allow_html=True)
            st.number_input("Jan 1 Price (¢/bu)", value=float(sn_jan1), step=0.25, key=f"{_P}sn_jan1", format="%.2f")
            sn_co = st.number_input("World C/O % of Use", value=float(sn_co_raw), step=0.1, key=f"{_P}sn_co", format="%.1f")
            st.selectbox("Current Year Category", ["No Crop Scare", "SA Crop Problem", "Crop Scare"], key=f"{_P}sn_section")
            st.markdown(f'<span class="hl-ratio-badge">World C/O: {sn_co:.2f}%</span>', unsafe_allow_html=True)

        st.markdown(f'<div style="margin-top:24px;color:{_MUTED};font-size:0.72rem;text-align:center;">'
                    'Data: USDA NASS &nbsp;·&nbsp; Model by JPSI &nbsp;·&nbsp; Changes take effect immediately.</div>',
                    unsafe_allow_html=True)

    # ══ CZ TAB ════════════════════════════════════════════════════════════════
    with tab_cz:
        st.markdown('<div class="hl-sec-hdr">Dec Corn (CZ) — Production as % of Same Year\'s Use</div>', unsafe_allow_html=True)
        m1, m2, m3, m4 = st.columns(4)
        m1.metric("Jan 1 Price",    f"${cz_jan1/100:.2f}/bu")
        m2.metric("US Production",  f"{cz_prod:,.0f} Mil Bu")
        m3.metric("US Usage",       f"{cz_use:,.0f} Mil Bu")
        m4.metric("Prod/Use Ratio", f"{cz_ratio*100:.2f}%",
                  delta="Surplus" if cz_ratio >= 1 else "Deficit",
                  delta_color="normal" if cz_ratio >= 1 else "inverse")
        _render_contract_tab(
            "CZ — Dec Corn", D["cz_low"], D["cz_high"],
            "ratio", "Prod/Use %", cz_ratio, cz_jan1,
            cz_hl_low, cz_lp, cz_t5_low, cz_hl_high, cz_hp, cz_t5_high,
            "Jan–Dec Low (¢/bu)", "Jan–Dec High (¢/bu)",
            "Dec Corn: Production/Use vs Jan–Dec H/L as % of Jan 1",
            "Dec Corn — Month When Annual High or Low Typically Occurs",
            current_section=cz_section,
            tile_low_label="Jan–Dec Indicated Low", tile_high_label="Jan–Dec Indicated High",
        )

    # ══ CN TAB ════════════════════════════════════════════════════════════════
    with tab_cn:
        st.markdown('<div class="hl-sec-hdr">Jul Corn (CN) — Carryout as % of Use</div>', unsafe_allow_html=True)
        m1, m2 = st.columns(2)
        m1.metric("Jan 1 Price",   f"${cn_jan1/100:.2f}/bu")
        m2.metric("Carryout/Use",  f"{cn_co_pct*100:.2f}%")
        _render_contract_tab(
            "CN — Jul Corn", D["cn_low"], D["cn_high"],
            "carryout_pct", "Carryout % of Use", cn_co_pct, cn_jan1,
            cn_hl_low, cn_lp, cn_t5_low, cn_hl_high, cn_hp, cn_t5_high,
            "Jan–Jul Low (¢/bu)", "Jan–Jul High (¢/bu)",
            "Jul Corn: Carryout/Use vs Jan–Jul H/L as % of Jan 1",
            "Jul Corn — Month When Annual High or Low Typically Occurs",
            tile_low_label="Jan–Jul Indicated Low", tile_high_label="Jan–Jul Indicated High",
        )

    # ══ SX TAB ════════════════════════════════════════════════════════════════
    with tab_sx:
        st.markdown('<div class="hl-sec-hdr">Nov Soybeans (SX) — World Carryout as % of Same Year\'s Use</div>', unsafe_allow_html=True)
        m1, m2 = st.columns(2)
        m1.metric("Jan 1 Price",     f"${sx_jan1/100:.2f}/bu")
        m2.metric("World C/O/Use",   f"{sx_co_pct*100:.2f}%")
        _render_contract_tab(
            "SX — Nov Beans", D["sx_low"], D["sx_high"],
            "carryout_pct", "World C/O % of Use", sx_co_pct, sx_jan1,
            sx_hl_low, sx_lp, sx_t5_low, sx_hl_high, sx_hp, sx_t5_high,
            "Jan–Nov Low (¢/bu)", "Jan–Nov High (¢/bu)",
            "Nov Soybeans: World C/O/Use vs Jan–Nov H/L as % of Jan 1",
            "Nov Soybeans — Month When Annual High or Low Typically Occurs",
            current_section=sx_section,
            tile_low_label="Jan–Nov Indicated Low", tile_high_label="Jan–Nov Indicated High",
        )

    # ══ SN TAB ════════════════════════════════════════════════════════════════
    with tab_sn:
        st.markdown(
            f'<div class="hl-sec-hdr">Jul Soybeans (SN) — World Carryout as % of Use '
            f'<span style="color:{_MUTED};font-size:0.75rem;font-weight:400;">'
            f'· {_SECTION_EMOJI.get(sn_section,"⚪")} {sn_section}</span></div>',
            unsafe_allow_html=True,
        )
        m1, m2, m3 = st.columns(3)
        m1.metric("Jan 1 Price",      f"${sn_jan1/100:.2f}/bu")
        m2.metric("World C/O/Use",    f"{sn_co_pct*100:.2f}%")
        m3.metric("Current Category", sn_section)
        _render_contract_tab(
            "SN — Jul Beans", D["sn_low"], D["sn_high"],
            "carryout_pct", "World C/O % of Use", sn_co_pct, sn_jan1,
            sn_hl_low, sn_lp, sn_t5_low, sn_hl_high, sn_hp, sn_t5_high,
            "Jan–Jul Low (¢/bu)", "Jan–Jul High (¢/bu)",
            "Jul Soybeans: World C/O/Use vs Jan–Jul H/L as % of Jan 1",
            "Jul Soybeans — Month When Annual High or Low Typically Occurs",
            current_section=sn_section,
            tile_low_label="Jan–Jul Indicated Low", tile_high_label="Jan–Jul Indicated High",
        )

    # ══ SEASONALS TAB ═════════════════════════════════════════════════════════
    with tab_seas:
        st.markdown('<div class="hl-sec-hdr">📈 Seasonal Price Patterns with Indicated H/L Levels</div>',
                    unsafe_allow_html=True)

        # Indicated H/L summary tiles
        st.markdown('<div class="hl-row-label">📉 Indicated Lows</div>', unsafe_allow_html=True)
        t1, t2, t3, t4 = st.columns(4)
        t1.markdown(_headline_tile("CZ — Dec Corn",  "Jan–Dec",  cz_hl_low,  cz_lp,  None, "low"),  unsafe_allow_html=True)
        t2.markdown(_headline_tile("CN — Jul Corn",  "Jan–Jul",  cn_hl_low,  cn_lp,  None, "low"),  unsafe_allow_html=True)
        t3.markdown(_headline_tile("SX — Nov Beans", "Jan–Nov",  sx_hl_low,  sx_lp,  None, "low"),  unsafe_allow_html=True)
        t4.markdown(_headline_tile("SN — Jul Beans", "Jan–Jul",  sn_hl_low,  sn_lp,  None, "low"),  unsafe_allow_html=True)
        st.markdown('<div style="height:6px;"></div>', unsafe_allow_html=True)
        st.markdown('<div class="hl-row-label">📈 Indicated Highs</div>', unsafe_allow_html=True)
        t1, t2, t3, t4 = st.columns(4)
        t1.markdown(_headline_tile("CZ — Dec Corn",  "Jan–Dec",  cz_hl_high, cz_hp, None, "high"), unsafe_allow_html=True)
        t2.markdown(_headline_tile("CN — Jul Corn",  "Jan–Jul",  cn_hl_high, cn_hp, None, "high"), unsafe_allow_html=True)
        t3.markdown(_headline_tile("SX — Nov Beans", "Jan–Nov",  sx_hl_high, sx_hp, None, "high"), unsafe_allow_html=True)
        t4.markdown(_headline_tile("SN — Jul Beans", "Jan–Jul",  sn_hl_high, sn_hp, None, "high"), unsafe_allow_html=True)
        st.markdown('<div style="height:14px;"></div>', unsafe_allow_html=True)

        # Contract selector
        seas_contract = st.radio(
            "Contract",
            ["🌽 Dec Corn (CZ26)", "🌽 Jul Corn (CN26)",
             "🫘 Nov Soybeans (SX26)", "🫘 Jul Soybeans (SN26)"],
            horizontal=True, key=f"{_P}seas_contract",
        )

        SEAS_CFG = {
            "🌽 Dec Corn (CZ26)": dict(
                product_code="ZC", month_letter="Z", fallback_letter=None, fallback_end_month=None,
                end_month=12, jan1_price=cz_jan1, current_year=2026,
                ind_high_price=_to_dollar(cz_hl_high), ind_low_price=_to_dollar(cz_hl_low),
                title="Dec Corn (CZ26) — Seasonal Price Pattern",
            ),
            "🌽 Jul Corn (CN26)": dict(
                product_code="ZC", month_letter="N", fallback_letter="Z", fallback_end_month=7,
                end_month=7, jan1_price=cn_jan1, current_year=2026,
                ind_high_price=_to_dollar(cn_hl_high), ind_low_price=_to_dollar(cn_hl_low),
                title="Jul Corn (CN26) — Seasonal Price Pattern",
            ),
            "🫘 Nov Soybeans (SX26)": dict(
                product_code="ZS", month_letter="X", fallback_letter=None, fallback_end_month=None,
                end_month=11, jan1_price=sx_jan1, current_year=2026,
                ind_high_price=_to_dollar(sx_hl_high), ind_low_price=_to_dollar(sx_hl_low),
                title="Nov Soybeans (SX26) — Seasonal Price Pattern",
            ),
            "🫘 Jul Soybeans (SN26)": dict(
                product_code="ZS", month_letter="N", fallback_letter="Z", fallback_end_month=7,
                end_month=7, jan1_price=sn_jan1, current_year=2026,
                ind_high_price=_to_dollar(sn_hl_high), ind_low_price=_to_dollar(sn_hl_low),
                title="Jul Soybeans (SN26) — Seasonal Price Pattern",
            ),
        }

        cfg = SEAS_CFG[seas_contract]

        with st.spinner("Loading price history…"):
            contracts_by_year = _build_seasonal_contracts(
                cfg["product_code"], cfg["month_letter"], api_key, legacy_df,
                cfg["current_year"],
                fallback_letter=cfg.get("fallback_letter"),
                fallback_end_month=cfg.get("fallback_end_month"),
            )

        seas_df = build_seasonal_df_hl(contracts_by_year, end_month=cfg["end_month"])

        if seas_df.empty:
            st.warning("No price history loaded for this contract. Check your Massive API key.")
        else:
            hist_df = seas_df[seas_df["year"] != cfg["current_year"]]
            curr_df = seas_df[seas_df["year"] == cfg["current_year"]]
            n_yrs   = hist_df["year"].nunique()
            _j      = cfg["jan1_price"] / 100 / 100
            avg_hi  = hist_df.groupby("year")["price_pct"].max().mean() * _j
            avg_lo  = hist_df.groupby("year")["price_pct"].min().mean() * _j
            curr_last = curr_df["close_raw"].iloc[-1] if not curr_df.empty else None

            mc1, mc2, mc3, mc4 = st.columns(4)
            mc1.metric("Historical Years",    str(n_yrs))
            mc2.metric("Avg Seasonal High",   f"${avg_hi:.2f}/bu")
            mc3.metric("Avg Seasonal Low",    f"${avg_lo:.2f}/bu")
            mc4.metric("Current Yr (Latest)", f"${curr_last/100:.2f}/bu" if curr_last else "No data")

            st.markdown('<div style="height:10px;"></div>', unsafe_allow_html=True)

            # Chart 1: clean averages
            st.markdown(
                f'<div class="hl-sec-div">KEY AVERAGES — {cfg["title"].split("—")[0].strip()}</div>',
                unsafe_allow_html=True,
            )
            st.plotly_chart(
                _make_seasonal_overlay(
                    seas_df, cfg["current_year"], cfg["title"] + "  ·  Averages & Current Year",
                    end_month=cfg["end_month"],
                    ind_high_price=cfg["ind_high_price"], ind_low_price=cfg["ind_low_price"],
                    jan1_price=cfg["jan1_price"], show_individual_years=False, height=540,
                ),
                use_container_width=True, key=f"hl_seas_clean_{seas_contract}",
            )

            # Chart 2: all historical years
            st.markdown(
                f'<div class="hl-sec-div">ALL HISTORICAL YEARS — {cfg["title"].split("—")[0].strip()}</div>',
                unsafe_allow_html=True,
            )
            st.plotly_chart(
                _make_seasonal_overlay(
                    seas_df, cfg["current_year"], cfg["title"],
                    end_month=cfg["end_month"],
                    ind_high_price=cfg["ind_high_price"], ind_low_price=cfg["ind_low_price"],
                    jan1_price=cfg["jan1_price"], show_individual_years=True, height=500,
                ),
                use_container_width=True, key=f"hl_seas_full_{seas_contract}",
            )

            with st.expander("📋 Year-by-Year Summary Table", expanded=False):
                _j2 = cfg["jan1_price"] / 100 / 100
                rows = []
                for yr, g in seas_df.groupby("year"):
                    j1_raw = g[g["doy"] == g["doy"].min()]["close_raw"].iloc[0]
                    rows.append({
                        "Year":          int(yr),
                        "Jan 1 ($/bu)":  f"${j1_raw/100:.2f}",
                        "Seasonal High": f"${g['price_pct'].max() * _j2:.2f}",
                        "Seasonal Low":  f"${g['price_pct'].min() * _j2:.2f}",
                        "Range":         f"${(g['price_pct'].max() - g['price_pct'].min()) * _j2:.2f}",
                        "Current Yr":    "★" if yr == cfg["current_year"] else "",
                    })
                st.dataframe(
                    pd.DataFrame(rows).sort_values("Year", ascending=False),
                    use_container_width=True, hide_index=True,
                )
